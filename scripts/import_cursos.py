#!/usr/bin/env python3
"""Importa cursos do Excel + CSV de preços para o Postgres site_anhanguera."""

from __future__ import annotations

import csv
import json
import os
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

import psycopg2
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = Path(r"C:\Users\Caio\Downloads")
SCHEMA = Path(__file__).with_name("schema.sql")

FEATURED_SLUGS = {"administracao", "pedagogia", "farmacia", "nutricao"}

NAME_ALIASES = {
    "ciencias da computacao": "ciencia da computacao",
    "computacao na nuvem": "computacao em nuvem",
    "engenharia da computacao": "engenharia de computacao",
    "gestao do terceiro setor": "gestao das organizacoes do terceiro setor",
    "gestao de producao industrial": "gestao da producao industrial",
    "gestao de qualidade": "gestao da qualidade",
    "gestao de tecnologia da informacao": "gestao da tecnologia da informacao",
    "seguranca do trabalho": "seguranca no trabalho",
    "terapias integrativas complementares": "terapias integrativas e complementares",
}

GRAD_CATEGORY = {
    "gestao e negocios": "Negócios",
    "engenharia e tecnologia": "Tecnologia",
    "arquitetura, design": "Arquitetura",
    "arquitetura design": "Arquitetura",
    "educacao": "Educação",
    "saude": "Saúde",
    "direito": "Jurídico",
    "comunicacao": "Comunicação",
    "gastronomia": "Gastronomia",
}

POS_CATEGORY = {
    "comunicacao": "Comunicação",
    "direito": "Jurídico",
    "educacao": "Educação",
    "engenharia": "Tecnologia",
    "gestao e negocios": "Negócios",
    "saude": "Saúde",
    "tecnologia": "Tecnologia",
}

FALLBACK_IMAGE = (
    "https://images.unsplash.com/photo-1523240795612-9a054b0db644?auto=format&fit=crop&w=800&q=80"
)


def load_dotenv(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip()
        if (value.startswith('"') and value.endswith('"')) or (
            value.startswith("'") and value.endswith("'")
        ):
            value = value[1:-1]
        os.environ.setdefault(key.strip(), value)


def norm(value: object) -> str:
    text = str(value or "").replace("\n", " ").strip().lower()
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn"
    )
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def slugify(value: object) -> str:
    text = norm(value).replace(" ", "-")
    return re.sub(r"-{2,}", "-", text).strip("-")


def parse_money(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip()
    if not text or text.upper() in {"#N/D", "N/D", "NONE"}:
        return None
    text = text.replace("R$", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        digits = re.sub(r"[^\d.]", "", text)
        try:
            return Decimal(digits) if digits else None
        except InvalidOperation:
            return None


def parse_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return int(value)
    digits = re.sub(r"[^\d]", "", str(value))
    return int(digits) if digits else None


def split_title_body(value: object) -> dict[str, str] | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip()
    if ": " in text:
        title, body = text.split(": ", 1)
        return {"title": title.strip(), "description": body.strip()}
    return {"title": text, "description": ""}


def items_from_row(row: tuple, indexes: list[int]) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    for idx in indexes:
        parsed = split_title_body(row[idx] if idx < len(row) else None)
        if parsed:
            items.append(parsed)
    return items


def bullets(value: object) -> list[dict[str, str]]:
    if not value:
        return []
    items: list[dict[str, str]] = []
    for line in str(value).splitlines():
        line = re.sub(r"^[\s\-•*]+", "", line).strip()
        if not line:
            continue
        if line.lower().startswith("entre as possibilidades"):
            continue
        items.append({"title": line})
    return items


def is_image_url(value: object) -> bool:
    text = str(value or "").strip()
    return text.startswith("http://") or text.startswith("https://")


def map_category(raw: object, mapping: dict[str, str], fallback: str) -> str:
    key = norm(raw)
    return mapping.get(key, fallback)


def parse_chave(chave: object, modalidade_col: object) -> tuple[str | None, str]:
    text = str(chave or "")
    lowered = text.lower()
    formacao = None
    if "licenciatura" in lowered:
        formacao = "Licenciatura"
    elif "bacharelado" in lowered:
        formacao = "Bacharelado"
    elif "tecn" in lowered:
        formacao = "Tecnólogo"

    modalidade = str(modalidade_col or "").strip()
    if not modalidade:
        if "semipresencial" in lowered:
            modalidade = "Semipresencial"
        else:
            modalidade = "EAD"
    if modalidade not in {"EAD", "Semipresencial"}:
        modalidade = "EAD" if "ead" in modalidade.lower() else modalidade
    return formacao, modalidade


def find_xlsx() -> Path:
    for path in DOWNLOADS.glob("BD - Eduit*.xlsx"):
        if "completa" in path.name.lower() or "completa" in path.name:
            return path
    matches = list(DOWNLOADS.glob("BD - Eduit*.xlsx"))
    if not matches:
        raise FileNotFoundError("Excel de cursos não encontrado em Downloads")
    return max(matches, key=lambda p: p.stat().st_mtime)


def find_csv() -> Path:
    matches = list(DOWNLOADS.glob("preço*.csv")) + list(DOWNLOADS.glob("preco*.csv"))
    if not matches:
        matches = [
            p
            for p in DOWNLOADS.glob("*.csv")
            if "preco" in p.name.lower() or "preço" in p.name.lower() or "preco+" in p.name.lower()
        ]
    if not matches:
        raise FileNotFoundError("CSV de preços não encontrado em Downloads")
    return matches[0]


def find_wix_graduacao() -> Path:
    candidates = [DOWNLOADS / "graduacao.csv", ROOT / "graduacao.csv"]
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError("CSV do CMS Wix (graduacao.csv) não encontrado em Downloads")


def csv_get(row: dict[str, str], *names: str) -> str:
    folded = {re.sub(r"\s+", " ", key).strip().lower(): key for key in row}
    for name in names:
        key = folded.get(name.strip().lower())
        if key and str(row.get(key) or "").strip():
            return str(row[key]).strip()
    return ""


def normalize_duracao(value: object) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"(\d+)\s*semestres?", text, re.I)
    if match:
        return f"{match.group(1)} semestres"
    return text


def load_wix_graduacao(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [row for row in csv.DictReader(handle) if csv_get(row, "Curso")]


def wix_grad_fields(row: dict[str, str]) -> dict:
    title = csv_get(row, "Curso")
    classificacao = csv_get(row, "Classificação")
    categoria = map_category(classificacao, GRAD_CATEGORY, classificacao or "Outros")
    semi = csv_get(row, "Semipresencial").lower() == "sim"
    chave = csv_get(row, "Chave 1")
    _, chave_mod = parse_chave(chave, "Semipresencial" if semi else "")
    duracao = normalize_duracao(csv_get(row, "Duração"))
    url_path = csv_get(row, "Graduacao- cursos")
    if url_path and not url_path.startswith("/"):
        url_path = f"/{url_path}"
    image = csv_get(row, "Image")
    area = csv_get(row, "area de atuação")
    return {
        "titulo": title,
        "categoria": categoria,
        "categoria_raw": classificacao,
        "category_badge": categoria.upper(),
        "sobre": csv_get(row, "Descrição") or None,
        "imagem": image if is_image_url(image) else None,
        "url_path": url_path or None,
        "formacao": csv_get(row, "Formação") or None,
        "duracao": duracao,
        "modalidade": "Semipresencial" if semi else chave_mod or "EAD",
        "meta_title": csv_get(row, "META TITLLE", "META TITLE") or None,
        "meta_description": csv_get(row, "META DESCRIPTION") or None,
        "img_alt": csv_get(row, "IMG alt text") or title,
        "mercado_trabalho": csv_get(row, "Mercado de Trabalho") or None,
        "area_atuacao_texto": area or None,
        "areas_atuacao": json.dumps(bullets(area), ensure_ascii=False),
    }


def sync_wix_graduacao(cur, rows: list[dict[str, str]]) -> tuple[int, list[str]]:
    cur.execute("SELECT id, titulo FROM cursos WHERE tipo = 'graduacao'")
    by_title = {norm(title): course_id for course_id, title in cur.fetchall()}
    updated = 0
    missing: list[str] = []
    for row in rows:
        fields = wix_grad_fields(row)
        key = NAME_ALIASES.get(norm(fields["titulo"]), norm(fields["titulo"]))
        course_id = by_title.get(key) or by_title.get(norm(fields["titulo"]))
        if course_id is None:
            missing.append(fields["titulo"])
            continue
        cur.execute(
            """
            UPDATE cursos SET
              categoria = %(categoria)s,
              categoria_raw = %(categoria_raw)s,
              category_badge = %(category_badge)s,
              sobre = COALESCE(%(sobre)s, sobre),
              imagem = COALESCE(%(imagem)s, imagem),
              url_path = COALESCE(%(url_path)s, url_path),
              formacao = COALESCE(%(formacao)s, formacao),
              duracao = %(duracao)s,
              modalidade = %(modalidade)s,
              meta_title = COALESCE(%(meta_title)s, meta_title),
              meta_description = COALESCE(%(meta_description)s, meta_description),
              img_alt = COALESCE(%(img_alt)s, img_alt),
              mercado_trabalho = COALESCE(%(mercado_trabalho)s, mercado_trabalho),
              area_atuacao_texto = COALESCE(%(area_atuacao_texto)s, area_atuacao_texto),
              areas_atuacao = CASE
                WHEN %(areas_atuacao)s = '[]' THEN areas_atuacao
                ELSE %(areas_atuacao)s::jsonb
              END
            WHERE id = %(id)s
            """,
            {**fields, "id": course_id},
        )
        if fields["duracao"]:
            cur.execute(
                "UPDATE curso_ofertas SET duracao = %s WHERE curso_id = %s",
                (fields["duracao"], course_id),
            )
        updated += 1
    return updated, missing


def load_prices(csv_path: Path) -> dict[str, list[dict]]:
    by_name: dict[str, list[dict]] = {}
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            title = (row.get("Curso") or "").replace("\n", " ").strip()
            if not title:
                continue
            offers: list[dict] = []
            groups = [
                (row.get("Chave 1"), row.get("preço 1") or row.get("preco1"), row.get("Modalidade 1")),
                (row.get("Chave 2"), row.get("preço 2 ") or row.get("preço 2"), row.get("Modalidade 2")),
                (row.get("chave 3 ") or row.get("chave 3"), row.get("preço 3 ") or row.get("preço 3"), row.get("Modalidade 3")),
            ]
            for chave, preco, modalidade in groups:
                valor = parse_money(preco)
                if valor is None and not str(chave or "").strip():
                    continue
                formacao, modalidade_ok = parse_chave(chave, modalidade)
                if valor is None:
                    continue
                offers.append(
                    {
                        "chave": str(chave or "").strip() or None,
                        "valor": valor,
                        "modalidade": modalidade_ok,
                        "formacao": formacao,
                    }
                )
            key = NAME_ALIASES.get(norm(title), norm(title))
            by_name[key] = offers
            by_name[norm(title)] = offers
    return by_name


def match_offers(title: str, prices: dict[str, list[dict]]) -> list[dict]:
    key = NAME_ALIASES.get(norm(title), norm(title))
    if key in prices:
        return prices[key]
    for price_key, offers in prices.items():
        if key in price_key or price_key in key:
            return offers
    return []


def db_connect():
    load_dotenv(ROOT / ".env")
    return psycopg2.connect(
        host=os.environ["DATABASE_HOST"],
        port=int(os.environ.get("DATABASE_PORT") or 5432),
        user=os.environ["DATABASE_USER"],
        password=os.environ["DATABASE_PASSWORD"],
        dbname=os.environ["DATABASE_NAME"],
        sslmode=os.environ.get("DATABASE_SSL") or "disable",
        connect_timeout=15,
    )


def insert_course(cur, payload: dict, offers: list[dict]) -> None:
    cur.execute(
        """
        INSERT INTO cursos (
            slug, tipo, titulo, categoria, categoria_raw, category_badge, sobre, imagem,
            url_path, formacao, duracao, modalidade, preco, meta_title, meta_description,
            img_alt, mercado_trabalho, area_atuacao_texto, indicacao_texto, aprendizados,
            indicacoes, areas_atuacao, featured
        ) VALUES (
            %(slug)s, %(tipo)s, %(titulo)s, %(categoria)s, %(categoria_raw)s, %(category_badge)s,
            %(sobre)s, %(imagem)s, %(url_path)s, %(formacao)s, %(duracao)s, %(modalidade)s,
            %(preco)s, %(meta_title)s, %(meta_description)s, %(img_alt)s, %(mercado_trabalho)s,
            %(area_atuacao_texto)s, %(indicacao_texto)s, %(aprendizados)s, %(indicacoes)s,
            %(areas_atuacao)s, %(featured)s
        )
        RETURNING id
        """,
        payload,
    )
    curso_id = cur.fetchone()[0]
    for offer in offers:
        cur.execute(
            """
            INSERT INTO curso_ofertas (curso_id, chave, duracao, parcelas, modalidade, formacao, valor)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (
                curso_id,
                offer.get("chave"),
                offer.get("duracao"),
                offer.get("parcelas"),
                offer.get("modalidade"),
                offer.get("formacao"),
                offer.get("valor"),
            ),
        )


def unique_slug(used: set[str], base: str) -> str:
    slug = base or "curso"
    n = 2
    while slug in used:
        slug = f"{base}-{n}"
        n += 1
    used.add(slug)
    return slug


def import_graduacao(ws, prices: dict[str, list[dict]], cur, used_slugs: set[str]) -> tuple[int, int]:
    imported = 0
    unmatched = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        title = str(row[0] or "").strip() if row else ""
        if not title:
            continue
        url_path = str(row[15] or "").strip()
        slug = unique_slug(used_slugs, slugify(url_path.split("/")[-1] if url_path else title))
        offers = match_offers(title, prices)
        if not offers:
            unmatched += 1
            print(f"SEM PREÇO (grad): {title}")
        primary = offers[0] if offers else {}
        formacao = primary.get("formacao")
        modalidade = primary.get("modalidade") or "EAD"
        duracao = None
        for offer in offers:
            if not offer.get("duracao"):
                offer["duracao"] = duracao
        categoria_raw = str(row[16] or "").strip()
        categoria = map_category(categoria_raw, GRAD_CATEGORY, categoria_raw or "Outros")
        payload = {
            "slug": slug,
            "tipo": "graduacao",
            "titulo": title,
            "categoria": categoria,
            "categoria_raw": categoria_raw,
            "category_badge": categoria.upper(),
            "sobre": str(row[1]).strip() if row[1] else None,
            "imagem": str(row[14]).strip() if is_image_url(row[14]) else FALLBACK_IMAGE,
            "url_path": url_path or f"/graduacao-cruzeiro/{slug}",
            "formacao": formacao,
            "duracao": duracao,
            "modalidade": modalidade,
            "preco": primary.get("valor"),
            "meta_title": str(row[18]).strip() if row[18] else None,
            "meta_description": str(row[19]).strip() if row[19] else None,
            "img_alt": str(row[17]).strip() if row[17] else title,
            "mercado_trabalho": None,
            "area_atuacao_texto": None,
            "indicacao_texto": None,
            "aprendizados": json.dumps(items_from_row(row, [2, 3, 4, 5]), ensure_ascii=False),
            "indicacoes": json.dumps(items_from_row(row, [6, 7, 8, 9]), ensure_ascii=False),
            "areas_atuacao": json.dumps(items_from_row(row, [10, 11, 12, 13]), ensure_ascii=False),
            "featured": slug in FEATURED_SLUGS,
        }
        insert_course(cur, payload, offers)
        imported += 1
    return imported, unmatched


def import_pos(ws, cur, used_slugs: set[str]) -> int:
    imported = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = str(row[0] or "").strip() if row else ""
        if not title:
            continue
        slug = unique_slug(used_slugs, slugify(title))
        categoria_raw = str(row[17] or "").strip()
        categoria = map_category(categoria_raw, POS_CATEGORY, categoria_raw or "Outros")
        offers: list[dict] = []
        primary = {
            "chave": str(row[11] or "").strip() or None,
            "duracao": str(row[2] or "").strip() or None,
            "parcelas": parse_int(row[3]),
            "modalidade": "EAD",
            "formacao": None,
            "valor": parse_money(row[5]),
        }
        if primary["valor"] is not None or primary["duracao"]:
            offers.append(primary)
        alt = {
            "chave": str(row[12] or "").strip() or None,
            "duracao": str(row[6] or "").strip() or None,
            "parcelas": parse_int(row[7]),
            "modalidade": "EAD",
            "formacao": None,
            "valor": parse_money(row[9]),
        }
        if alt["duracao"] or alt["valor"] is not None:
            offers.append(alt)
        chosen = next((o for o in offers if o.get("valor") is not None), offers[0] if offers else {})
        payload = {
            "slug": slug,
            "tipo": "pos",
            "titulo": title,
            "categoria": categoria,
            "categoria_raw": categoria_raw,
            "category_badge": "PÓS-GRADUAÇÃO",
            "sobre": str(row[13]).strip() if row[13] else None,
            "imagem": str(row[18]).strip() if is_image_url(row[18]) else FALLBACK_IMAGE,
            "url_path": f"/pos/{slug}",
            "formacao": "Especialização" if "mba" not in title.lower() else "MBA",
            "duracao": chosen.get("duracao") or "12 Meses",
            "modalidade": "EAD",
            "preco": chosen.get("valor"),
            "meta_title": None,
            "meta_description": None,
            "img_alt": title,
            "mercado_trabalho": str(row[14]).strip() if row[14] else None,
            "area_atuacao_texto": str(row[15]).strip() if row[15] else None,
            "indicacao_texto": str(row[16]).strip() if row[16] else None,
            "aprendizados": json.dumps([], ensure_ascii=False),
            "indicacoes": json.dumps([], ensure_ascii=False),
            "areas_atuacao": json.dumps(bullets(row[15]), ensure_ascii=False),
            "featured": False,
        }
        if "mba" in title.lower():
            payload["formacao"] = "MBA"
        insert_course(cur, payload, offers)
        imported += 1
    return imported


def main() -> None:
    if "--sync-wix-grad" in sys.argv:
        wix_path = find_wix_graduacao()
        print("Wix CMS:", wix_path)
        rows = load_wix_graduacao(wix_path)
        conn = db_connect()
        conn.autocommit = False
        try:
            with conn.cursor() as cur:
                updated, missing = sync_wix_graduacao(cur, rows)
            conn.commit()
            print(f"Graduação atualizada do CMS: {updated}/{len(rows)}")
            if missing:
                print("Sem match no banco:")
                for title in missing:
                    print(f"  - {title}")
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
        return

    xlsx = find_xlsx()
    csv_path = find_csv()
    print("Excel:", xlsx)
    print("CSV:", csv_path)
    prices = load_prices(csv_path)
    wb = load_workbook(xlsx, data_only=True)
    conn = db_connect()
    conn.autocommit = False
    try:
        with conn.cursor() as cur:
            cur.execute(SCHEMA.read_text(encoding="utf-8"))
            cur.execute("TRUNCATE curso_ofertas, cursos RESTART IDENTITY")
            used: set[str] = set()
            grad_n, unmatched = import_graduacao(wb[wb.sheetnames[0]], prices, cur, used)
            pos_n = import_pos(wb[wb.sheetnames[1]], cur, used)
            cur.execute("SELECT COUNT(*) FROM cursos")
            total = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM curso_ofertas")
            offers = cur.fetchone()[0]
        conn.commit()
        print(f"Graduação: {grad_n} (sem preço: {unmatched})")
        print(f"Pós: {pos_n}")
        print(f"Total cursos: {total} | ofertas: {offers}")
    except Exception:
        conn.rollback()
        raise
    finally:
        wb.close()
        conn.close()


if __name__ == "__main__":
    main()
